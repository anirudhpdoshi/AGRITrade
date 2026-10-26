"""Coverage round 2 (read-only): CONAB xlsx layout, CONAB crop-progress page, Comex Stat exports API."""
import io, json, os, re, datetime as dt
import requests
OUT = os.path.join(os.path.dirname(__file__), "..", "coverage"); os.makedirs(OUT, exist_ok=True)
UA = {"User-Agent": "Mozilla/5.0 (desk coverage check)"}
R = {"run_at": dt.datetime.utcnow().isoformat() + "Z", "conab_xlsx": {}, "progress": {}, "comex": {}}
B = "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/safra-de-graos/boletim-da-safra-de-graos/"
X = {"out-2026 (1o 2026/27)": B + "1o-levantamento-safra-2026-27/site_previsao_de_safra-por_produto-out-2026.xlsx",
     "ago-2026 (11o 2025/26)": B + "11o-levantamento-safra-2025-26/site_previsao_de_safra-por_produto-ago-2026.xlsx"}
import openpyxl
for k, u in X.items():
    try:
        r = requests.get(u, headers=UA, timeout=60); d = {"status": r.status_code, "bytes": len(r.content)}
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
        d["sheets"] = wb.sheetnames
        for sn in wb.sheetnames:
            if re.search(r"milho|soja|trigo|total", sn, re.I):
                ws = wb[sn]; rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 45: break
                    rows.append([("" if v is None else str(v))[:28] for v in row[:14]])
                d["sheet_" + sn] = rows
        R["conab_xlsx"][k] = d
    except Exception as e:
        R["conab_xlsx"][k] = {"error": str(e)[:300]}
try:
    u = "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/progresso-de-safra/acompanhamento-das-lavouras-12-10-a-18-10-26"
    r = requests.get(u, headers=UA, timeout=60); R["progress"]["status"] = r.status_code
    R["progress"]["links"] = sorted(set(re.findall(r'href="([^"]+\.(?:xlsx?|pdf|csv)[^"]*)"', r.text)))[:20]
    txt = re.sub(r"<[^>]+>", " ", r.text); txt = re.sub(r"\s+", " ", txt)
    i = txt.lower().find("soja"); R["progress"]["text_sample"] = txt[max(0, i - 300): i + 1500]
except Exception as e:
    R["progress"]["error"] = str(e)[:300]
C = "https://api-comexstat.mdic.gov.br"
try:
    R["comex"]["updated"] = requests.get(C + "/general/dates/updated", headers=UA, timeout=60).text[:300]
    body = {"flow": "export", "monthDetail": True, "period": {"from": "2026-01", "to": "2026-12"},
            "filters": [{"filter": "ncm", "values": ["10059010", "12019000", "10019900"]}],
            "details": ["ncm", "country"], "metrics": ["metricFOB", "metricKG"]}
    r = requests.post(C + "/general?language=en", json=body, headers=UA, timeout=120)
    R["comex"]["status"] = r.status_code
    js = r.json(); R["comex"]["keys"] = list(js.keys()) if isinstance(js, dict) else type(js).__name__
    lst = (js.get("data") or {}).get("list") if isinstance(js, dict) and isinstance(js.get("data"), dict) else js.get("data") if isinstance(js, dict) else js
    R["comex"]["n"] = len(lst or []); R["comex"]["sample"] = (lst or [])[:6]
except Exception as e:
    R["comex"]["error"] = str(e)[:300]
open(os.path.join(OUT, "coverage2.json"), "w").write(json.dumps(R, indent=1, default=str))
