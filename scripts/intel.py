"""Global Commodity Intelligence: official institutional data + event headlines, each mapped to what it means for corn,
soybeans and wheat on this desk.

Writes intel.json to the orphan branch `intel` (read by intel-ui.js, the "Intel" tab). Separate job like world.py and
global_grains.py: if it fails, nothing else on the desk is touched.

Sources (all keyless public):
  World Bank  Pink Sheet monthly prices (CMO-Historical-Data-Monthly.xlsx) + Commodity Markets Outlook publication watch
  IMF         Primary Commodity Price System (PCPS) via the `imfp` client - BACKUP for any World Bank series that fails,
              and an independent cross-check where both report the same month
  FAO         Food Price Index + Cereals / Vegetable-oils sub-indices (monthly CSV) and next release date
  Headlines   Google News RSS: World Bank, IMF, FAO, OECD-FAO, AMIS, plus supply-disruption / trade / energy / fertilizer

Rules (same as the rest of the desk):
  * Every block carries source, source URL, fetched-at time and the month the data runs through.
  * A failed source keeps its last good copy, flagged stale, with the original fetch time - never silently refreshed.
  * Surprises are z-scores against the series' own history; revisions only against a value we stored earlier.
  * Directional bias is RULE-BASED and labelled as inference. Official figures are 'verified'; headlines are not.
    Nothing here is a trading signal: every item lists what must be confirmed first.
"""
import csv
import datetime as dt
import hashlib
import io
import json
import math
import os
import re
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
from email.utils import parsedate_to_datetime
from urllib.parse import quote_plus, urlparse, urljoin

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
WT = "/tmp/intel-worktree"
BRANCH = "intel"
UA = {"User-Agent": "Mozilla/5.0 (grain desk intelligence feed; contact via repository)"}
CROPS = ("corn", "soybeans", "wheat")

WB_PAGE = "https://www.worldbank.org/en/research/commodity-markets"
WB_XLSX_FALLBACK = [
    "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx",
    "https://thedocs.worldbank.org/en/doc/18675f1d1639c7a34d463f59263ba0a2-0050012025/related/CMO-Historical-Data-Monthly.xlsx",
]
FAO_PAGE = "https://www.fao.org/worldfoodsituation/foodpricesindex/en/"
FAO_CSV_FALLBACK = "https://www.fao.org/media/docs/worldfoodsituationlibraries/wfs-library/food_price_indices_data.csv"
IMF_PAGE = "https://data.imf.org/Datasets/PCPS"

# ------------------------------------------------------------------ indicator catalogue + transmission map
# wb: Pink Sheet column header prefix (case-insensitive). imf: PCPS indicator code. fao: FAO CSV column.
# eff: per-crop effect of a RISE in this series: +1 direct supportive, +.5 indirect supportive, -.5 indirect pressure.
# exp: which desk exposures it touches (flat price, basis, spreads, vol, FX, margins).
INDICATORS = [
    {"id": "maize", "name": "Maize (US Gulf)", "cat": "ag", "unit": "$/mt", "wb": "maize", "imf": "PMAIZMT",
     "eff": {"corn": 1, "wheat": .5}, "hz": "confirming", "exp": ["flat price"],
     "why": "World export price of the crop itself. A monthly average, so it confirms what CBOT already did rather than leading it; "
            "a gap between this and the current futures level shows how far the market has moved since the month closed. Wheat follows "
            "through feed substitution."},
    {"id": "wheat", "name": "Wheat (US HRW)", "cat": "ag", "unit": "$/mt", "wb": "wheat, us hrw", "imf": "PWHEAMT",
     "eff": {"wheat": 1, "corn": .5}, "hz": "confirming", "exp": ["flat price", "spreads"],
     "why": "World wheat export price (US hard red winter). Confirms the month's price level; competing origins (Black Sea, EU, "
            "Argentina, Australia) decide whether US and Canadian wheat stay competitive. Corn follows through feed substitution."},
    {"id": "soybeans", "name": "Soybeans", "cat": "ag", "unit": "$/mt", "wb": "soybeans", "imf": "PSOYB",
     "eff": {"soybeans": 1}, "hz": "confirming", "exp": ["flat price", "basis"],
     "why": "World soybean price. Confirms the month; the Brazil-versus-US export spread decides where China buys, which drives "
            "US and Canadian export basis."},
    {"id": "soy_oil", "name": "Soybean oil", "cat": "ag", "unit": "$/mt", "wb": "soybean oil", "imf": "PSOIL",
     "eff": {"soybeans": 1}, "hz": "1–3 months", "exp": ["flat price", "margins"],
     "why": "Higher soybean oil lifts the crush margin, so crushers bid harder for beans. Biodiesel and renewable-diesel policy is the "
            "main demand driver; watch the board crush and soybean-oil share."},
    {"id": "palm_oil", "name": "Palm oil", "cat": "ag", "unit": "$/mt", "wb": "palm oil", "imf": "PPOIL",
     "eff": {"soybeans": .5}, "hz": "1–3 months", "exp": ["margins"],
     "why": "Palm oil is the main substitute for soybean oil. Dearer palm oil pulls import demand toward soybean oil, which supports "
            "crush margins and so soybean demand."},
    {"id": "urea", "name": "Urea (fertilizer)", "cat": "energy_fert", "unit": "$/mt", "wb": "urea", "imf": "PUREA",
     "eff": {"corn": .5, "soybeans": -.5, "wheat": .5}, "hz": "next season", "exp": ["new-crop spreads", "margins"],
     "why": "Nitrogen is corn's largest variable cost and soybeans need almost none. Dearer urea squeezes corn margins and can shift "
            "next season's acres from corn to soybeans: supportive for new-crop corn, a weight on new-crop soybeans. Winter wheat "
            "faces the same input cost. Slow channel: it shows up in acreage, not in this week's price."},
    {"id": "dap", "name": "DAP (fertilizer)", "cat": "energy_fert", "unit": "$/mt", "wb": "dap", "imf": None,
     "eff": {"corn": .5, "wheat": .5}, "hz": "next season", "exp": ["new-crop spreads", "margins"],
     "why": "Phosphate cost. Raises the cost of growing corn and wheat; affordability can cut application rates and yields at the margin."},
    {"id": "crude", "name": "Crude oil (avg)", "cat": "energy_fert", "unit": "$/bbl", "wb": "crude oil, average", "imf": "POILAPSP",
     "eff": {"corn": .5, "soybeans": .5}, "hz": "weeks", "exp": ["flat price", "basis", "FX"],
     "why": "Higher crude improves ethanol (corn) and biodiesel (soybean oil) economics. It also raises freight and drying costs, "
            "which widens interior basis, and it tends to strengthen the CAD, which lowers the CAD flat price for the same CBOT price."},
    {"id": "natgas_eu", "name": "Natural gas (Europe)", "cat": "energy_fert", "unit": "$/mmbtu", "wb": "natural gas, europe", "imf": "PNGASEU",
     "eff": {"corn": .5, "wheat": .5}, "hz": "next season", "exp": ["margins"],
     "why": "Natural gas is the feedstock for nitrogen fertilizer. A sustained rise in European gas cuts nitrogen output and lifts "
            "world fertilizer prices a few months later."},
    {"id": "fao_cereals", "name": "FAO Cereals index", "cat": "ag", "unit": "index 2014-16=100", "fao": "cereals",
     "eff": {"wheat": 1, "corn": .5}, "hz": "confirming", "exp": ["flat price"],
     "why": "FAO's export-quotation basket for wheat, maize, rice and barley. Confirms whether world grain import costs rose or fell "
            "last month."},
    {"id": "fao_oils", "name": "FAO Vegetable-oils index", "cat": "ag", "unit": "index 2014-16=100", "fao": "oils",
     "eff": {"soybeans": .5}, "hz": "confirming", "exp": ["margins"],
     "why": "World vegetable-oil prices (palm, soy, sunflower, rapeseed). Rising oils support crush margins and so soybean demand."},
    {"id": "fao_ffpi", "name": "FAO Food Price Index", "cat": "macro_trade", "unit": "index 2014-16=100", "fao": "food price index",
     "eff": {}, "hz": "context", "exp": ["policy risk"],
     "why": "Headline world food inflation. It does not move CBOT directly, but high or rising food prices raise the odds of export "
            "restrictions by producer governments, which is a supply-disruption risk for wheat and corn."},
    {"id": "imf_food", "name": "IMF Food price index", "cat": "macro_trade", "unit": "index", "imf": "PFOOD", "imf_only": True,
     "eff": {}, "hz": "context", "exp": ["policy risk"],
     "why": "IMF's broad food basket. Macro context for food inflation in importing countries and commodity-linked currencies."},
]
IND = {i["id"]: i for i in INDICATORS}
NEUTRAL_PCT = 0.5          # a monthly move smaller than this is 'no material change'
MATERIAL_Z = 1.5           # |z| at or above this is a material surprise against the series' own history
DIVERGE_PCT = 3.0          # World Bank vs IMF disagreement on the same month that gets flagged

CATS = {"ag": "Agriculture fundamentals", "energy_fert": "Energy & fertilizers", "macro_trade": "Macro & trade",
        "disruption": "Supply disruption", "outlook": "Long-term outlook"}

# ------------------------------------------------------------------ headline queries (Google News RSS)
QUERIES = [
    ("worldbank", "official", '"World Bank" (commodity OR "food prices" OR fertilizer OR "Pink Sheet" OR "Commodity Markets Outlook")'),
    ("imf", "official", 'IMF ("commodity prices" OR "food prices" OR "food inflation" OR "commodity exporters")'),
    ("fao", "official", 'FAO ("Food Price Index" OR cereal OR "crop prospects" OR "Food Outlook" OR "cereal supply")'),
    ("oecd", "official", '"OECD-FAO" OR "Agricultural Outlook" OECD'),
    ("amis", "official", '"AMIS Market Monitor" OR "Agricultural Market Information System"'),
    ("disrupt", "event", '(wheat OR corn OR soybean OR grain) ("export ban" OR "export restriction" OR "export quota" OR "export tax" OR drought OR frost OR flood)'),
    ("logistics", "event", '(grain OR wheat OR corn OR soybean) (port OR "Black Sea" OR "Panama Canal" OR "Mississippi River" OR barge OR freight OR strike)'),
    ("fert", "event", '(fertilizer OR urea OR potash OR phosphate) prices (farmers OR supply OR export)'),
    ("trade", "event", '(soybean OR wheat OR corn) (tariff OR "trade deal" OR China purchases OR devaluation OR real OR peso)'),
]
OFFICIAL_DOM = ("worldbank.org", "imf.org", "fao.org", "oecd.org", "usda.gov", "amis-outlook.org", "agr.gc.ca", "statcan.gc.ca", "conab.gov.br")

KW_CROP = {
    "corn": r"\b(corn|maize|ethanol)\b",
    "soybeans": r"\b(soy\w*|crush\w*|biodiesel|renewable diesel)\b",
    "wheat": r"\bwheat\b",
}
KW_OTHER = {"rice": r"\brice\b", "veg oils": r"\b(palm oil|vegetable oils?|sunflower|rapeseed|canola)\b",
            "fertilizer": r"\b(fertili[sz]er|urea|potash|phosphate|nitrogen|ammonia)\b",
            "energy": r"\b(crude|oil prices?|brent|wti|diesel|natural gas|energy)\b", "freight": r"\b(freight|shipping|barge|vessel|port|canal)\b"}
KW_CAT = [
    ("disruption", r"\b(drought|flood\w*|frost|heat ?wave|crop failure|pest|disease|war|attack\w*|strike|port clos\w*|blockade|canal|river level|low water|export ban|export restriction\w*|halt\w*|missile\w*|drone\w*|strait|hormuz|deadl\w*|shipping route|suspen\w*|el ni[nñ]o|la ni[nñ]a)\b"),
    ("energy_fert", r"\b(fertili[sz]er|urea|potash|phosphate|ammonia|natural gas|crude|oil price|diesel|biofuel|ethanol|biodiesel|freight)\b"),
    ("macro_trade", r"\b(tariff\w*|trade (?:deal|war|talks)|export (?:tax|quota)|currency|devalu\w*|dollar|interest rate|inflation|growth|gdp|recession|sanction\w*)\b"),
    ("outlook", r"\b(outlook|projection\w*|long[- ]term|medium[- ]term|2030|2035|decade)\b"),
    ("ag", r"\b(harvest|crop|planting|acre\w*|yield|production|stocks|inventor\w*|exports?|imports?|cereal|grain)\b"),
]
REGIONS = [("Black Sea", r"\b(black sea|ukrain\w*|russia\w*|odesa|novorossiysk)\b"), ("Brazil", r"\b(brazil\w*|conab|mato grosso|santos|paranagua)\b"),
           ("Argentina", r"\b(argentin\w*|rosario|buenos aires)\b"), ("US", r"\b(u\.s\.|us|usda|united states|american|midwest|corn belt|mississippi|gulf)\b"),
           ("Canada", r"\b(canad\w*|prairie\w*|saskatchewan|alberta|manitoba|ontario|vancouver)\b"), ("China", r"\b(china|chinese|beijing|sinograin)\b"),
           ("India", r"\b(india\w*)\b"), ("EU", r"\b(eu|europe\w*|france|french|germany|euronext)\b"), ("Australia", r"\b(australia\w*)\b"),
           ("Middle East / N. Africa", r"\b(egypt\w*|morocco|algeria|saudi|iran|red sea|suez)\b")]
# grain-price direction read from headline text (supply-side words; 'cut forecast' = less supply = supportive)
BULL = [r"export ban", r"export restriction\w*", r"restrict\w* exports", r"export (?:tax|quota)",
        r"drought (?:\w+ ){0,2}(?:hit\w*|hurt\w*|damag\w*|cut\w*|threat\w*|stress\w*|worsen\w*|grip\w*|parch\w*|slash\w*|wither\w*|loss\w*)",
        r"(?:hit|hurt|damaged|cut|parched|scorched|threatened|withered|ravaged) by (?:\w+ )?drought", r"drought-(?:hit|stricken|damaged|ravaged)",
        r"frost (?:\w+ ){0,2}(?:damag\w*|hit\w*|threat\w*)", r"heat ?wave", r"crop failure",
        r"(?:cut|cuts|lowers|lowered|lowering|reduc\w*|slash\w*|downgrad\w*) (?:(?!as\b|on\b|amid\b)\w+ ){0,2}(?:forecast|estimate|output|production|crop|harvest)", r"(?:output|production|harvest|crop) (?:falls?|drops?|declines?|shrinks?)",
        r"smallest", r"lowest stocks", r"tight\w* (?:\w+ )?(?:supply|supplies|stocks)", r"port (?:closure|closed|attack\w*)", r"attack\w* on (?:port|grain)", r"halt\w* (?:exports|shipments)",
        r"(?:prices?|futures|market) (?:\w+ ){0,3}(?:rise|rises|rose|jump\w*|surge\w*|soar\w*|climb\w*|spike\w*|rall\w*|leap\w*)", r"prices? up\b", r"\bleap in (?:\w+ ){0,2}prices",
        r"hit\w* (?:new |fresh |record |\d+-\w+ )?highs?", r"strong (?:demand|buying|imports)", r"china (?:buys|purchases|books)", r"higher (?:prices|costs)",
        r"suspen\w* (?:\w+ )?(?:production|output|exports|shipments)", r"supply (?:\w+ )?(?:restrict\w*|disrupt\w*|shortage\w*|squeez\w*)", r"shortage\w*", r"risk premium (?:\w+ )?(?:ris\w*|build\w*|return\w*)",
        r"increase\w* (?:\w+ ){0,3}(?:purchases|buying|imports)", r"bounce\w*", r"china buying", r"(?:plantings?|acreage|area|sowings?) (?:\w+ ){0,3}(?:\w+-\w+ )?low",
        r"high (?:\w+ )?prices (?:\w+ ){0,2}linger", r"(?:harvest|crop|yield) loss\w*", r"(?:production|output|harvest|crop|yield\w*) (?:\w+ ){0,3}(?:near |to |at )?(?:a )?(?:historic|record|multi-year|\d+-year) low",
        r"shortens? (?:\w+ )?(?:crops?|harvest)", r"drought (?:\w+ ){0,1}(?:pushes|drives|sends|shortens|reduces|reshap\w*)"]
BEAR = [r"record (?:\w+ ){0,2}(?:crop|harvest|output|production|exports?)", r"(?:drift\w*|trad\w*|edg\w*|slip\w*|eas\w*|sink\w*|mov\w*|turn\w*|clos\w*|end\w*|open\w*|start\w*|finish\w*) lower", r"(?:futures|prices?|grains?|corn|wheat|soybeans?|beans) (?:are |were |)lower\b",
        r"(?:grains?|corn|wheat|soybeans?|beans|soy) (?:(?!imports?\b|exports?\b|plantings?\b|acreage\b|output\b|production\b|stocks?\b)\w+ ){0,3}(?:fall|falls|fell|slip\w*|drop\w*|sink\w*|slump\w*|tumbl\w*|retreat\w*)", r"bumper", r"(?:raise\w*|lift\w*|increase\w*|upgrad\w*|boost\w*) (?:\w+ ){0,2}(?:forecast|estimate|output|production|crop)",
        r"(?:output|production|harvest|crop) (?:rises?|jumps?|grows?|climbs?)", r"ample (?:supply|supplies|stocks)", r"(?:lift\w*|end\w*|scrap\w*|remov\w*) (?:the )?export ban",
        r"resum\w* (?:exports|shipments)", r"glut", r"surplus",
        r"(?:prices?|futures|market) (?:\w+ ){0,3}(?:fall|falls|fell|drop\w*|slump\w*|declin\w*|eas\w*|slid\w*|tumbl\w*|soften\w*|weaken\w*|sag\w*|six-year low|lowest)", r"prices? down\b",
        r"hit\w* (?:new |fresh |multi-\w+ |\d+-\w+ )?lows?", r"weak (?:demand|buying|imports)", r"weaker (?:currency|real|peso)", r"pressures? (?:\w+ ){0,3}prices",
        r"cancel\w* (?:purchases|cargoes|orders)", r"lower (?:prices|costs)", r"increase\w* in (?:\w+ ){0,3}(?:production|output|forecast)", r"risk premium (?:\w+ )?(?:deflat\w*|eas\w*|fad\w*|evaporat\w*)",
        r"supplies stabili[sz]e\w*", r"(?:harvest|crop|yield\w*|output|production) (?:\w+ ){0,3}(?:exceed\w*|beat\w*|top\w*|above) (?:\w+ ){0,2}expectations",
        r"better[- ]than[- ]expected (?:\w+ ){0,2}(?:harvest|crop|yield\w*|output)", r"imports? (?:\w+ ){0,3}(?:could |may |to |will |set to )?(?:fall|drop|declin\w*|plunge\w*|slump\w*|shrink\w*)", r"(?:production|output|crop|harvest) (?:\w+ ){0,3}(?:to rise|to grow|seen higher)",
        r"(?:prices?|grain|grains|futures|wheat|corn|soybeans?) (?:\w+ ){0,3}\w+-(?:week|month|year) low", r"(?:lift|lifts|lifted|end|ends|ended) (?:\w+ ){0,2}ban"]
# a headline must mention a commodity, price, food or farm-economy term to enter the feed (filters off-topic matches)
RELEVANT = r"\b(commodit\w*|food|price\w*|grain\w*|cereal\w*|wheat|corn|maize|soy\w*|fertili[sz]er\w*|urea|potash|phosphate|crude|oil|energy|crop\w*|harvest\w*|agricultur\w*|outlook|inflation|freight|shipping|export\w*|import\w*|tariff\w*|drought|planting|yield\w*|biofuel\w*|ethanol)\b"
PER_QUERY = 25             # newest headlines kept per query per run

CONFIRM = {
    "corn": "CBOT corn curve (Z26/H27 spread), WASDE balance sheet, weekly export sales, Ontario/US Gulf basis",
    "soybeans": "CBOT soybeans and board crush (oil share), Brazil–US FOB spread, weekly export sales to China",
    "wheat": "CBOT/KC/MGEX wheat spreads, Black Sea FOB offers, Canadian Grain Commission exports, Egypt/Algeria tenders",
}


# ------------------------------------------------------------------ small helpers (pure, unit-tested)
def now_utc():
    return dt.datetime.now(dt.timezone.utc)


def pct(cur, prev):
    return None if cur is None or prev in (None, 0) else round(100.0 * (cur - prev) / prev, 2)


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if isinstance(v, float) and math.isnan(v) else float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "…", "..", "-", "–", "n.a.", "na"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def month_key(v):
    """'2026M09' / '2026-09' / '2026-09-01' / datetime / '9/1/2026' -> '2026-09'; None if not a month."""
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v or "").strip()
    m = re.match(r"^(\d{4})\s*[M\-/.](\d{1,2})(?:\D|$)", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return f"{int(m.group(3)):04d}-{int(m.group(1)):02d}"
    return None


def prev_month(k):
    y, m = int(k[:4]), int(k[5:])
    return f"{y - 1:04d}-12" if m == 1 else f"{y:04d}-{m - 1:02d}"


def stats_move(series):
    """Latest monthly move of a {'YYYY-MM': value} series, with a z-score of that % move against the previous 24 monthly % moves.
    Returns None when there are fewer than two months."""
    ks = sorted(k for k, v in series.items() if v is not None)
    if len(ks) < 2:
        return None
    last, prev = ks[-1], ks[-2]
    v, p = series[last], series[prev]
    chg_pct = pct(v, p)
    hist = []
    for a, b in zip(ks[-26:-2], ks[-25:-1]):
        c = pct(series[b], series[a])
        if c is not None:
            hist.append(c)
    z = None
    if len(hist) >= 12 and chg_pct is not None:
        mu = sum(hist) / len(hist)
        sd = math.sqrt(sum((x - mu) ** 2 for x in hist) / (len(hist) - 1))
        z = round((chg_pct - mu) / sd, 2) if sd > 0 else None
    yk = f"{int(last[:4]) - 1:04d}{last[4:]}"
    return {"month": last, "value": v, "prev_month": prev, "prev": p, "chg": round(v - p, 4), "chg_pct": chg_pct,
            "yoy_pct": pct(v, series.get(yk)), "z": z, "material": bool(z is not None and abs(z) >= MATERIAL_Z)}


def revisions(series, stored, tol_pct=0.05):
    """Months whose value changed versus the snapshot we stored on an earlier run (never invented: no snapshot -> [])."""
    out = []
    for k in sorted(series):
        a, b = stored.get(k), series[k]
        if a is None or b is None or a == 0:
            continue
        if abs(100.0 * (b - a) / a) > tol_pct:
            out.append({"month": k, "prev": a, "new": b, "pct": pct(b, a)})
    return out[-3:]


def bias_word(score):
    return "bullish" if score > 0 else "bearish" if score < 0 else "neutral"


def impact_from_move(ind, mv):
    """Rule-based per-crop read of an official series' latest monthly move. Returns {crop: {...}} and overall label."""
    out = {}
    if not mv or mv.get("chg_pct") is None:
        return out, "neutral", "Low"
    flat = abs(mv["chg_pct"]) < NEUTRAL_PCT
    sgn = 0 if flat else (1 if mv["chg_pct"] > 0 else -1)
    for crop in CROPS:
        e = ind.get("eff", {}).get(crop)
        if e is None:
            continue
        s = sgn * e
        strength = "direct" if abs(e) >= 1 else "indirect"
        out[crop] = {"bias": "neutral" if flat else bias_word(s), "strength": strength, "horizon": ind["hz"]}
    signs = {v["bias"] for v in out.values() if v["bias"] != "neutral"}
    overall = "neutral" if not signs else signs.pop() if len(signs) == 1 else "mixed"
    direct = any(v["strength"] == "direct" for v in out.values())
    conf = "High" if (direct and mv.get("material")) else "Medium" if (direct or mv.get("material")) and not flat else "Low"
    return out, overall, conf


def norm_title(t):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", (t or "").lower())).strip()


def tokens(t):
    stop = {"the", "a", "an", "of", "to", "in", "on", "for", "and", "as", "at", "by", "with", "from", "is", "are", "its", "after", "amid", "over", "says", "say"}
    return {w for w in norm_title(t).split() if len(w) > 2 and w not in stop}


def similar(a, b, thr=0.6):
    A, B = tokens(a), tokens(b)
    return bool(A and B) and len(A & B) / len(A | B) >= thr


def classify_headline(title):
    """Category, crops, other commodities, region(s), directional read and the phrases that produced it (the evidence)."""
    low = (title or "").lower()
    crops = [c for c, rx in KW_CROP.items() if re.search(rx, low)]
    broad = None
    if not crops and re.search(r"\bcereals?\b", low):
        crops, broad = ["corn", "wheat"], "cereals"
    elif not crops and re.search(r"\bgrains?\b", low):
        crops, broad = list(CROPS), "grains"
    other = [c for c, rx in KW_OTHER.items() if re.search(rx, low)]
    cat = next((c for c, rx in KW_CAT if re.search(rx, low)), "ag")
    regions = [r for r, rx in REGIONS if re.search(rx, low)] or ["Global"]
    bull = [m.group(0) for rx in BULL for m in [re.search(rx, low)] if m]
    bull = [b for b in bull if not re.search(r"\b(?:despite|in spite of|despite the|shrugs? off|weathers?|tolerance|resistan\w*|insurance|buffer|relief)\b(?:\W+\w+){0,3}\W+" + re.escape(b.split()[0]), low)
            and not re.search(re.escape(b.split()[-1]) + r"\W+(?:\w+\W+){0,1}(?:tolerance|resistan\w*|insurance|relief|buffer|eases?|easing|ends?|over)\b", low)]   # 'despite drought', 'drought tolerance' are not supply cuts
    # a price rise that belongs to inputs (fertilizer, farm costs) is a cost story, not a grain-price signal
    if re.search(r"\b(?:fertili[sz]er|input|farm|production) (?:prices?|costs?)\b", low):
        bull = [b for b in bull if not re.search(r"\b(?:prices?|costs?)\b", b)]
    bull = [b for b in bull if not (b == "bounce" and re.search(r"\b(?:try|tries|trying|attempt\w*|or bottom)\b", low))]   # 'try to bounce?' is not a bounce
    bear = [m.group(0) for rx in BEAR for m in [re.search(rx, low)] if m]
    bear = [x for x in bear if not re.search(r"planting|acreage|sowing|\barea\b", x)]   # 'plantings at a 7-year low' is a supply cut, not a price low
    score = len(bull) - len(bear)
    # an 'export ban lifted' style phrase must beat the bare 'export ban' match it contains
    if any("export ban" in b for b in bear) and "export ban" in bull:
        score -= 1
    return {"cat": cat, "crops": crops, "broad": broad, "other": other, "regions": regions,
            "bias": "mixed" if bull and bear and score == 0 else bias_word(score), "bull": bull, "bear": bear}


def headline_impact(cls):
    """Per-crop read for a headline. Named crops get the headline's direction; input-cost headlines map through the fertilizer /
    energy channel; anything else stays 'watch'. Always inference from text only."""
    out = {}
    b = cls["bias"] if cls["bias"] in ("bullish", "bearish", "mixed") else None
    targets = cls["crops"] or []
    if not targets and "fertilizer" in cls["other"]:
        for crop, eff in (("corn", .5), ("wheat", .5), ("soybeans", -.5)):
            if b in ("bullish", "bearish"):
                # a rise in fertilizer prices is a cost (supportive new-crop corn/wheat, weight on new-crop soy)
                sg = (1 if b == "bullish" else -1) * (1 if eff > 0 else -1)
                out[crop] = {"bias": bias_word(sg), "strength": "indirect", "horizon": "next season"}
        return out
    if not targets and ("energy" in cls["other"] or "freight" in cls["other"]):
        for crop in ("corn", "soybeans"):
            out[crop] = {"bias": "watch", "strength": "indirect", "horizon": "weeks"}
        return out
    for crop in targets:
        out[crop] = {"bias": b or "watch", "strength": "indirect" if cls.get("broad") else "direct", "horizon": "days–weeks" if cls["cat"] == "disruption" else "weeks"}
    return out


def overall_of(imp):
    s = {v["bias"] for v in imp.values() if v["bias"] in ("bullish", "bearish", "mixed")}
    return "watch" if not s else s.pop() if len(s) == 1 else "mixed"


def item_id(*parts):
    return hashlib.sha1("|".join(str(p) for p in parts).encode()).hexdigest()[:12]


def channels_for(cls):
    ch = []
    if cls["cat"] == "disruption":
        ch.append("supply")
    if "fertilizer" in cls["other"]:
        ch.append("fertilizer")
    if "energy" in cls["other"]:
        ch.append("energy / biofuels")
    if "freight" in cls["other"]:
        ch.append("freight / basis")
    if cls["cat"] == "macro_trade":
        ch.append("trade flows / FX")
    if not ch:
        ch.append("supply / demand")
    return ch


def first_weekday(y, m, wd):
    d = dt.date(y, m, 1)
    return d + dt.timedelta(days=(wd - d.weekday()) % 7)


def next_monthly(day_fn, today):
    d = day_fn(today.year, today.month)
    if d <= today:
        y, m = (today.year + 1, 1) if today.month == 12 else (today.year, today.month + 1)
        d = day_fn(y, m)
    return d


def bizday(y, m, n):
    """n-th business day (Mon–Fri) of a month."""
    d, c = dt.date(y, m, 1), 0
    while True:
        if d.weekday() < 5:
            c += 1
            if c == n:
                return d
        d += dt.timedelta(days=1)


# ------------------------------------------------------------------ fetchers
def http():
    import requests
    s = requests.Session()
    s.headers.update(UA)
    return s


def get(S, url, tries=3, **kw):
    kw.setdefault("timeout", 60)
    last = None
    for i in range(tries):
        try:
            r = S.get(url, **kw)
            if r.status_code == 429:
                time.sleep(10 * (i + 1))
                continue
            r.raise_for_status()
            return r
        except Exception as e:
            last = e
            time.sleep(2 * (i + 1))
    raise RuntimeError(str(last)[:200])


def find_links(html, pattern, base):
    import html as _h
    return [urljoin(base, _h.unescape(h)) for h in re.findall(r'href=["\']([^"\']+)["\']', html or "", re.I) if re.search(pattern, h, re.I)]


def parse_pinksheet(rows):
    """Rows of the 'Monthly Prices' sheet -> {indicator_id: {'YYYY-MM': value}}. The header row is the one that names 'Maize'."""
    hdr_i = next((i for i, r in enumerate(rows[:20]) if any(str(c or "").strip().lower() == "maize" for c in r)), None)
    if hdr_i is None:
        raise ValueError("Pink Sheet header row (Maize) not found")
    hdr = [str(c or "").strip().lower() for c in rows[hdr_i]]
    cols = {}
    for ind in INDICATORS:
        if not ind.get("wb"):
            continue
        j = next((j for j, h in enumerate(hdr) if h == ind["wb"]), None)
        if j is None:
            j = next((j for j, h in enumerate(hdr) if h.startswith(ind["wb"])), None)
        if j is not None:
            cols[ind["id"]] = j
    out = {k: {} for k in cols}
    for r in rows[hdr_i + 1:]:
        mk = month_key(r[0] if r else None)
        if not mk:
            continue
        for k, j in cols.items():
            v = num(r[j]) if j < len(r) else None
            if v is not None:
                out[k][mk] = v
    return {k: v for k, v in out.items() if v}


def worldbank_block(S):
    xlsx_url, page_ok, pubs = None, False, []
    try:
        html = get(S, WB_PAGE, tries=2).text
        page_ok = True
        links = find_links(html, r"CMO-Historical-Data-Monthly\.xlsx", WB_PAGE)
        xlsx_url = links[0] if links else None
        for u in find_links(html, r"CMO[-_ ]?(?:April|October)[-_ ]?\d{4}|commodity-markets-outlook", WB_PAGE):
            m = re.search(r"(April|October)[-_ ]?(\d{4})", u, re.I)
            if not m:
                continue
            pubs.append({"edition": f"{m.group(1)} {m.group(2)}", "url": u})
    except Exception:
        pass
    import openpyxl
    errs = []
    for u in ([xlsx_url] if xlsx_url else []) + WB_XLSX_FALLBACK:
        try:
            r = get(S, u, tries=2, timeout=120)
            wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
            sh = next((wb[n] for n in wb.sheetnames if n.strip().lower() == "monthly prices"), None)
            if sh is None:
                raise ValueError("sheet 'Monthly Prices' missing")
            series = parse_pinksheet([list(r) for r in sh.iter_rows(values_only=True)])
            if not series:
                raise ValueError("no series parsed")
            through = max(max(v) for v in series.values())
            uniq = {}
            for p in pubs:
                uniq.setdefault(p["edition"], p)
            pubs = sorted(uniq.values(), key=lambda p: (p["edition"][-4:], p["edition"].startswith("October")), reverse=True)
            series = {k: dict(sorted(v.items())[-60:]) for k, v in series.items()}
            return {"source": "World Bank Pink Sheet (CMO historical data, monthly)", "url": u, "page_ok": page_ok,
                    "fetched_at": now_utc().isoformat(), "through": through, "series": series, "cmo": pubs[:4]}
        except Exception as e:
            errs.append(f"{urlparse(u).path.split('/')[-1]}: {str(e)[:90]}")
    raise RuntimeError("; ".join(errs) or "no Pink Sheet URL")


def imf_block():
    """IMF PCPS via `imfp` (SDMX 3.0 client for the new data.imf.org API). Defensive: dimension and column names are discovered,
    not assumed; any code that returns nothing is reported, never filled."""
    import imfp
    want = {i["id"]: i["imf"] for i in INDICATORS if i.get("imf")}
    series, missing = {}, []
    for k, code in want.items():
        df = None
        for kw in ({"indicator": code, "frequency": "M", "data_transformation": "USD"},
                   {"indicator": code, "frequency": "M", "data_transformation": "INDEX"},
                   {"indicator": code, "frequency": "M"}):
            try:
                d = imfp.imf_get("PCPS", **kw)
                if d is not None and len(d):
                    df = d
                    break
            except Exception:
                continue
        if df is None:
            missing.append(code)
            continue
        cols = {c.lower(): c for c in df.columns}
        tc = cols.get("time_period") or cols.get("date") or cols.get("time")
        vc = cols.get("obs_value") or cols.get("value")
        if not tc or not vc:
            missing.append(code)
            continue
        s = {}
        for t, v in zip(df[tc], df[vc]):
            mk, x = month_key(t), num(v)
            if mk and x is not None:
                s[mk] = x
        if s:
            series[k] = dict(sorted(s.items())[-48:])
        else:
            missing.append(code)
        time.sleep(0.6)   # IMF asks for ~10 requests per 5 s
    if not series:
        raise RuntimeError("IMF PCPS returned no series (" + ", ".join(missing) + ")")
    return {"source": "IMF Primary Commodity Price System (PCPS)", "url": IMF_PAGE, "fetched_at": now_utc().isoformat(),
            "through": max(max(v) for v in series.values()), "series": series, "missing": missing}


def parse_fao_csv(text):
    rows = list(csv.reader(io.StringIO(text)))
    hdr_i = next((i for i, r in enumerate(rows[:15]) if any("cereal" in (c or "").lower() for c in r)), None)
    if hdr_i is None:
        raise ValueError("FAO header row (Cereals) not found")
    hdr = [(c or "").strip().lower() for c in rows[hdr_i]]
    cols = {}
    for ind in INDICATORS:
        if ind.get("fao"):
            j = next((j for j, h in enumerate(hdr) if h.startswith(ind["fao"])), None)
            if j is not None:
                cols[ind["id"]] = j
    out = {k: {} for k in cols}
    for r in rows[hdr_i + 1:]:
        mk = month_key(r[0] if r else None)
        if not mk:
            continue
        for k, j in cols.items():
            v = num(r[j]) if j < len(r) else None
            if v is not None:
                out[k][mk] = v
    return {k: dict(sorted(v.items())[-60:]) for k, v in out.items() if v}


def parse_fao_next(html):
    m = re.search(r"next release[^0-9A-Za-z]{0,20}(?:date)?[^0-9]{0,20}(\d{1,2}\s+[A-Z][a-z]+\s+\d{4})", html or "", re.I)
    if not m:
        return None
    try:
        return dt.datetime.strptime(m.group(1), "%d %B %Y").date().isoformat()
    except ValueError:
        return None


def fao_block(S):
    csv_url, nxt = None, None
    try:
        html = get(S, FAO_PAGE, tries=2).text
        links = find_links(html, r"food_price_indices_data[^\"']*\.csv", FAO_PAGE)
        csv_url = links[0] if links else None
        nxt = parse_fao_next(html)
    except Exception:
        pass
    errs = []
    for u in ([csv_url] if csv_url else []) + [FAO_CSV_FALLBACK]:
        try:
            r = get(S, u, tries=2)
            series = parse_fao_csv(r.content.decode("utf-8-sig", errors="replace"))
            if not series:
                raise ValueError("no series parsed")
            return {"source": "FAO Food Price Index (monthly)", "url": u, "fetched_at": now_utc().isoformat(),
                    "through": max(max(v) for v in series.values()), "series": series, "next_release": nxt}
        except Exception as e:
            errs.append(str(e)[:100])
    raise RuntimeError("; ".join(errs))


def rss(S, q, days=7):
    url = f"https://news.google.com/rss/search?q={quote_plus(q + f' when:{days}d')}&hl=en-CA&gl=CA&ceid=CA:en"
    root = ET.fromstring(get(S, url, tries=2, timeout=30).content)
    out = []
    for it in root.iter("item"):
        title, link = it.findtext("title") or "", it.findtext("link") or ""
        src = it.find("source")
        pub = src.text if src is not None and src.text else ""
        dom = urlparse(src.get("url")).netloc if src is not None and src.get("url") else "news.google.com"
        if pub and title.endswith(" - " + pub):
            title = title[: -len(" - " + pub)]
        try:
            when = parsedate_to_datetime(it.findtext("pubDate")).astimezone(dt.timezone.utc)
        except Exception:
            continue
        out.append({"ts": when.isoformat(), "title": title.strip(), "url": link, "outlet": pub or dom, "domain": dom.replace("www.", "")})
    out.sort(key=lambda r: r["ts"], reverse=True)
    return out[:PER_QUERY]


# ------------------------------------------------------------------ item builders
def data_items(merged, sources_used, snaps):
    """One feed item per official series' latest month. merged: {id: series}; sources_used: {id: 'World Bank'|'IMF (backup)'|'FAO'}."""
    items = []
    for k, s in merged.items():
        ind = IND[k]
        mv = stats_move(s)
        if not mv:
            continue
        imp, overall, conf = impact_from_move(ind, mv)
        revs = revisions(s, snaps.get(k, {}))
        src = sources_used[k]
        dirn = "unchanged" if abs(mv["chg_pct"] or 0) < NEUTRAL_PCT else ("up" if mv["chg_pct"] > 0 else "down")
        mname = dt.date(int(mv["month"][:4]), int(mv["month"][5:]), 1).strftime("%b %Y")
        title = f"{ind['name']} {dirn} {abs(mv['chg_pct'] or 0):.1f}% in {mname}" if dirn != "unchanged" else f"{ind['name']} unchanged in {mname} ({mv['chg_pct']:+.1f}%)"
        if mv["material"]:
            title += " — material move"
        items.append({
            "id": item_id("data", k, mv["month"], round(mv["value"], 4)), "kind": "data", "cat": ind["cat"], "ind": k,
            "ts": None, "title": title, "source": src, "verified": True, "material": mv["material"],
            "crops": [c for c in imp], "other": [], "regions": ["Global"],
            "key": {"label": ind["name"], "unit": ind["unit"], "month": mv["month"], "actual": mv["value"], "prev": mv["prev"],
                    "prev_month": mv["prev_month"], "chg": mv["chg"], "chg_pct": mv["chg_pct"], "yoy_pct": mv["yoy_pct"], "z": mv["z"]},
            "revisions": revs, "impact": imp, "bias": overall, "conf": conf, "horizon": ind["hz"], "exposure": ind["exp"],
            "why": ind["why"], "evidence": [f"{ind['name']}: {mv['prev']:,.2f} ({mv['prev_month']}) → {mv['value']:,.2f} ({mv['month']}) {ind['unit']}"]
                     + ([f"z-score {mv['z']:+.2f} vs the last 24 monthly moves"] if mv["z"] is not None else [])
                     + [f"revised {r['month']}: {r['prev']:,.2f} → {r['new']:,.2f}" for r in revs],
            "confirm": [CONFIRM[c] for c in imp] or ["Context series — no direct CBOT read"],
            "spark": [round(s[m], 3) for m in sorted(s)[-24:]],
        })
    return items


def headline_items(raw, prior_log):
    """Cluster near-duplicate headlines, classify, map to crops. Keeps the first time the desk saw each story."""
    seen_at = {i["id"]: i.get("seen") for i in prior_log}
    raw = sorted(raw, key=lambda r: r["ts"])
    clusters = []
    for r in raw:
        c = next((c for c in clusters if similar(c["title"], r["title"])), None)
        if c:
            c["also"].append(r["outlet"])
            if r["official"] and not c["official"]:
                c.update(url=r["url"], outlet=r["outlet"], domain=r["domain"], official=True)
            continue
        clusters.append(dict(r, also=[]))
    items, now = [], now_utc().isoformat()
    for c in clusters:
        if not re.search(RELEVANT, c["title"].lower()):
            continue   # off-topic (e.g. a World Bank infrastructure loan)
        cls = classify_headline(c["title"])
        if c["q"] in ("worldbank", "imf", "fao", "oecd", "amis") and cls["cat"] in ("ag", "macro_trade") and re.search(r"outlook|projection|long[- ]term", c["title"], re.I):
            cls["cat"] = "outlook"
        if not cls["crops"] and not cls["other"] and c["q"] in ("disrupt", "logistics", "trade"):
            continue   # off-topic match from a broad query
        imp = headline_impact(cls)
        iid = item_id("news", norm_title(c["title"]))
        offdom = any(c["domain"].endswith(d) for d in OFFICIAL_DOM)
        inst = {"worldbank": "World Bank", "imf": "IMF", "fao": "FAO", "oecd": "OECD–FAO", "amis": "AMIS"}.get(c["q"])
        conf = "Medium" if (offdom and (cls["bull"] or cls["bear"])) else "Low"
        alert = cls["cat"] == "disruption" and overall_of(imp) in ("bullish", "bearish") and bool(cls["crops"])
        items.append({
            "id": iid, "kind": "headline", "cat": cls["cat"], "ts": c["ts"], "seen": seen_at.get(iid) or now,
            "title": c["title"], "source": c["outlet"], "domain": c["domain"], "url": c["url"], "inst": inst,
            "verified": False, "official_src": offdom, "also": sorted(set(c["also"]))[:6],
            "crops": cls["crops"], "other": cls["other"], "regions": cls["regions"], "impact": imp, "bias": overall_of(imp),
            "conf": conf, "alert": alert, "channels": channels_for(cls),
            "evidence": [f"supportive phrase: “{p}”" for p in cls["bull"]] + [f"pressuring phrase: “{p}”" for p in cls["bear"]]
                        or ["no directional phrase in the headline — category and commodity only"],
            "confirm": [CONFIRM[c2] for c2 in (cls["crops"] or [])] or ["Read the source; check whether it changes a balance sheet, a cost or a trade flow"],
        })
    return items


def calendar(fao_next, today):
    """Next expected releases. Dates scraped from the publisher are marked 'published'; rule-of-thumb dates are 'estimated'."""
    cal = [
        {"src": "FAO", "what": "Food Price Index + Cereal Supply & Demand Brief",
         "date": fao_next or next_monthly(lambda y, m: first_weekday(y, m, 4), today).isoformat(), "basis": "published" if fao_next else "estimated (first Friday)"},
        {"src": "World Bank", "what": "Pink Sheet monthly prices",
         "date": next_monthly(lambda y, m: bizday(y, m, 2), today).isoformat(), "basis": "estimated (2nd business day)"},
        {"src": "IMF", "what": "Primary Commodity Prices (PCPS) monthly update",
         "date": next_monthly(lambda y, m: bizday(y, m, 4), today).isoformat(), "basis": "estimated (early month)"},
        {"src": "World Bank", "what": "Commodity Markets Outlook (semi-annual)",
         "date": (dt.date(today.year, 4, 28) if today < dt.date(today.year, 4, 20) else dt.date(today.year, 10, 28) if today < dt.date(today.year, 10, 20) else dt.date(today.year + 1, 4, 28)).isoformat(),
         "basis": "estimated (late April / late October)"},
        {"src": "OECD–FAO", "what": "Agricultural Outlook (10-year projections)",
         "date": (dt.date(today.year, 7, 8) if today < dt.date(today.year, 7, 1) else dt.date(today.year + 1, 7, 8)).isoformat(), "basis": "estimated (early July)"},
    ]
    return sorted(cal, key=lambda c: c["date"])


# ------------------------------------------------------------------ publish
def sh(*a, cwd=None, check=True):
    return subprocess.run(a, cwd=cwd, check=check, capture_output=True, text=True)


def load_branch_file(path):
    r = sh("git", "show", f"origin/{BRANCH}:{path}", cwd=ROOT, check=False)
    try:
        return json.loads(r.stdout) if r.returncode == 0 else None
    except Exception:
        return None


def publish(body, files):
    for k, v in (("user.name", "desk-intel-bot"), ("user.email", "41898282+github-actions[bot]@users.noreply.github.com")):
        sh("git", "config", k, v, cwd=ROOT)
    sh("rm", "-rf", WT, check=False)
    sh("git", "worktree", "prune", cwd=ROOT, check=False)
    exists = sh("git", "ls-remote", "--exit-code", "--heads", "origin", BRANCH, cwd=ROOT, check=False).returncode == 0
    if exists:
        sh("git", "worktree", "add", "-f", "-B", BRANCH, WT, f"origin/{BRANCH}", cwd=ROOT)
    else:
        sh("git", "worktree", "add", "--detach", WT, cwd=ROOT)
        sh("git", "checkout", "-q", "--orphan", BRANCH, cwd=WT)
        sh("git", "rm", "-rfq", ".", cwd=WT, check=False)
    os.makedirs(os.path.join(WT, "snap"), exist_ok=True)
    for name, data in files.items():
        with open(os.path.join(WT, name), "w", encoding="utf-8") as f:
            json.dump(data, f, separators=(",", ":"), ensure_ascii=False)
    with open(os.path.join(WT, "intel.json"), "w", encoding="utf-8") as f:
        json.dump(body, f, separators=(",", ":"), ensure_ascii=False)
    sh("git", "add", "-A", cwd=WT)
    sh("git", "commit", "-q", *([] if not exists else ["--amend"]), "-m", "intel: " + body["updated_at"], cwd=WT)
    sh("git", "push", "-q", "-f", "origin", BRANCH, cwd=WT)


def merge_series(wb, imf, fao):
    """Primary/backup per series: World Bank first, IMF fills any series World Bank lacks; FAO series are FAO's own.
    Where both World Bank and IMF report the latest common month, record the disagreement."""
    merged, used, checks = {}, {}, []
    wbs, imfs = (wb or {}).get("series", {}), (imf or {}).get("series", {})
    for ind in INDICATORS:
        k = ind["id"]
        if ind.get("fao"):
            if (fao or {}).get("series", {}).get(k):
                merged[k], used[k] = fao["series"][k], "FAO"
            continue
        if wbs.get(k):
            merged[k], used[k] = dict(sorted(wbs[k].items())[-48:]), "World Bank" + (" (last good copy)" if wb.get("stale") else "")
        elif imfs.get(k):
            merged[k], used[k] = imfs[k], "IMF (backup)" if not ind.get("imf_only") else "IMF"
        if wbs.get(k) and imfs.get(k):
            common = sorted(set(wbs[k]) & set(imfs[k]))
            if common:
                m = common[-1]
                d = pct(wbs[k][m], imfs[k][m])
                checks.append({"ind": k, "month": m, "wb": wbs[k][m], "imf": imfs[k][m], "diff_pct": d, "flag": d is not None and abs(d) > DIVERGE_PCT})
    return merged, used, checks


def main():
    t0 = time.time()
    S = http()
    sh("git", "fetch", "-q", "origin", BRANCH, cwd=ROOT, check=False)
    prev = load_branch_file("intel.json") or {}
    prev_raw = load_branch_file("snap/raw.json") or {}
    snaps = load_branch_file("snap/series.json") or {}
    log = (load_branch_file("snap/log.json") or {}).get("items", [])
    body = {"updated_at": now_utc().isoformat(), "schema": 1, "sources": {}, "health": {"ok": [], "failed": []}}

    blocks = {}
    for name, fn in (("worldbank", lambda: worldbank_block(S)), ("imf", imf_block), ("fao", lambda: fao_block(S))):
        try:
            blocks[name] = fn()
            body["health"]["ok"].append(name)
        except Exception as e:
            body["health"]["failed"].append({"part": name, "error": str(e)[:240]})
            p = prev_raw.get(name)
            if p:
                blocks[name] = dict(p, stale=True)
    raw_heads = []
    for q, kind, query in QUERIES:
        try:
            for r in rss(S, query):
                r.update(q=q, official=any(r["domain"].endswith(d) for d in OFFICIAL_DOM))
                raw_heads.append(r)
            body["health"]["ok"].append("news/" + q)
        except Exception as e:
            body["health"]["failed"].append({"part": "news/" + q, "error": str(e)[:160]})
        time.sleep(1)

    merged, used, checks = merge_series(blocks.get("worldbank"), blocks.get("imf"), blocks.get("fao"))
    items = data_items(merged, used, snaps)
    heads = headline_items(raw_heads, log) if raw_heads else [i for i in (prev.get("items") or []) if i.get("kind") == "headline"]
    if not raw_heads and heads:
        for i in heads:
            i["stale"] = True
    # data items carry the time the desk first saw that month's value
    for i in items:
        old = next((x for x in log if x["id"] == i["id"]), None)
        i["seen"] = (old or {}).get("seen") or body["updated_at"]
        i["ts"] = i["seen"]
    # searchable history: every item ever shown, newest first, capped
    by_id = {i["id"]: i for i in log}
    for i in items + heads:
        by_id[i["id"]] = i
    for i in by_id.values():   # stored headlines are re-read with the current classifier, so a fix applies to history too
        if i.get("kind") == "headline" and i.get("title"):
            try:
                cls = classify_headline(i["title"])
                imp = headline_impact(cls)
                i["impact"], i["bias"] = imp, overall_of(imp)
                i["evidence"] = ([f"supportive phrase: “{p}”" for p in cls["bull"]] + [f"pressuring phrase: “{p}”" for p in cls["bear"]]
                                 or ["no directional phrase in the headline — category and commodity only"])
                i["alert"] = cls["cat"] == "disruption" and i["bias"] in ("bullish", "bearish") and bool(cls["crops"])
            except Exception:
                pass
    cutoff = (now_utc() - dt.timedelta(days=45)).isoformat()
    hist = sorted(by_id.values(), key=lambda i: i.get("ts") or i.get("seen") or "", reverse=True)
    hist = [i for i in hist if (i.get("ts") or "") >= cutoff or i["kind"] == "data"][:400]

    for name, b in blocks.items():
        body["sources"][name] = {k: b.get(k) for k in ("source", "url", "fetched_at", "through", "stale", "next_release", "missing", "cmo", "page_ok") if b.get(k) is not None}
    body["sources"]["news"] = {"source": "Google News RSS (institutional + event queries)", "fetched_at": body["updated_at"] if raw_heads else (prev.get("sources", {}).get("news", {}).get("fetched_at")),
                               "stale": not raw_heads, "queries": len(QUERIES)}
    body["served_by"] = used
    body["crosscheck"] = checks
    body["calendar"] = calendar((blocks.get("fao") or {}).get("next_release"), now_utc().date())
    body["indicators"] = [{"id": k, "name": IND[k]["name"], "cat": IND[k]["cat"], "unit": IND[k]["unit"], "src": used[k], "move": stats_move(s),
                           "spark": [round(s[m], 3) for m in sorted(s)[-24:]], "months": sorted(s)[-24:]} for k, s in merged.items()]
    body["items"] = hist[:250]
    body["run_seconds"] = round(time.time() - t0, 1)

    new_snaps = {k: dict(sorted(v.items())[-36:]) for k, v in merged.items()}
    files = {"snap/series.json": new_snaps, "snap/log.json": {"items": hist}, "snap/raw.json": blocks}   # raw = last good copy per source
    if os.environ.get("INTEL_DRY"):
        out = os.environ.get("INTEL_OUT")
        if out:
            json.dump(body, open(out, "w"), indent=1, ensure_ascii=False)
        print(json.dumps({k: body[k] for k in ("updated_at", "health", "served_by", "crosscheck", "calendar")}, indent=1)[:4000])
        return
    publish(body, files)
    print(f"::notice title=intel::ok {len(body['health']['ok'])} · failed {len(body['health']['failed'])} · items {len(body['items'])} · {body['run_seconds']} s")
    for f in body["health"]["failed"]:
        print(f"::warning title=intel::{f['part']}: {f['error']}")


def diag_imf():
    """Print what the IMF PCPS API actually exposes (dimensions, a few codes, one query) so the series map can be checked."""
    import imfp
    print("imfp", getattr(imfp, "__version__", "?"))
    for fn, args in (("imf_get_datastructure", ("PCPS",)),):
        try:
            print(fn, "->", str(getattr(imfp, fn)(*args))[:1500])
        except Exception as e:
            print(fn, "FAILED", repr(e)[:300])
    try:
        cl = imfp.imf_get_codelists(["INDICATOR"], "PCPS")
        txt = str(cl)
        print("codelists ->", txt[:800])
        for w in ("MAIZ", "WHEA", "SOY", "UREA", "OIL", "FOOD"):
            print("  lines with", w, ":", [l for l in txt.splitlines() if w in l.upper()][:6])
    except Exception as e:
        print("codelists FAILED", repr(e)[:300])
    for kw in ({"indicator": "PMAIZMT", "frequency": "M"}, {"indicator": "PMAIZMT"}):
        try:
            d = imfp.imf_get("PCPS", **kw)
            print("imf_get", kw, "->", None if d is None else (list(d.columns), d.tail(3).to_dict("records")))
        except Exception as e:
            print("imf_get", kw, "FAILED", repr(e)[:300])


if __name__ == "__main__":
    if "--diag-imf" in sys.argv:
        diag_imf()
    else:
        main()
